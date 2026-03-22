import openpyxl

workbook = openpyxl.load_workbook(filename="C:/Users/Kundan_Kumar/Desktop/saucedemo_credentials.xlsx")
sheet = workbook['saucedemo']
username = sheet.cell(row=1,column=1).value
password =  sheet.cell(row=1,column=2).value


# y = [{'username': 'xyz'},{'password':'abc'}]
# print(y[0]['username'])